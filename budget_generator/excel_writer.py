"""Build a monthly budget .xlsx workbook with live formulas via openpyxl."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .currency import DEFAULT_CURRENCY, excel_number_format
from .layout import DEFAULT_EXPENSES, DEFAULT_INCOME, build_row_plan

TITLE_FONT = Font(size=16, bold=True, color="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(bold=True)
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
SUMMARY_FILL = PatternFill("solid", fgColor="E2EFDA")
SUMMARY_FONT = Font(bold=True, size=12)
THIN_BORDER = Border(bottom=Side(style="thin", color="CCCCCC"))
PERCENT_FORMAT = "0.0%"

STYLE_HANDLERS = {"title", "header", "section", "total", "summary"}


def build_workbook(
    month_label: str,
    incomes: list[tuple[str, float]] | None = None,
    expenses: list[tuple[str, float, float]] | None = None,
    currency: str = DEFAULT_CURRENCY,
) -> Workbook:
    incomes = incomes or DEFAULT_INCOME
    expenses = expenses or DEFAULT_EXPENSES
    plan = build_row_plan(month_label, incomes, expenses)
    currency_format = excel_number_format(currency)

    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "Budget"

    for row_idx, row_cells in enumerate(plan.rows, start=1):
        for col_idx, cell in enumerate(row_cells, start=1):
            xl_cell = ws.cell(row=row_idx, column=col_idx, value=cell.value)
            _apply_style(xl_cell, cell.style, col_idx)

    # Conditional formatting: highlight overspent categories (Difference < 0) in red.
    diff_range = f"D{plan.expense_start_row}:D{plan.expense_end_row}"
    ws.conditional_formatting.add(
        diff_range,
        FormulaRule(formula=[f"D{plan.expense_start_row}<0"], fill=PatternFill("solid", fgColor="FFC7CE")),
    )

    # Percent format for savings rate.
    ws.cell(row=plan.savings_rate_row, column=2).number_format = PERCENT_FORMAT

    # Currency formatting for every amount column across the sheet.
    for row_idx in range(1, len(plan.rows) + 1):
        for col_idx in (2, 3, 4):
            c = ws.cell(row=row_idx, column=col_idx)
            if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
                if row_idx != plan.savings_rate_row:
                    c.number_format = currency_format

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    for col_idx, width in enumerate((28, 16, 16, 16), start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    return wb


def _apply_style(xl_cell, style: str, col_idx: int) -> None:
    if style == "title":
        xl_cell.font = TITLE_FONT
        xl_cell.alignment = Alignment(horizontal="left")
    elif style == "section":
        xl_cell.font = SECTION_FONT
        xl_cell.fill = SECTION_FILL
    elif style == "header":
        xl_cell.font = HEADER_FONT
        xl_cell.fill = HEADER_FILL
        xl_cell.border = THIN_BORDER
    elif style == "total":
        xl_cell.font = TOTAL_FONT
        xl_cell.fill = TOTAL_FILL
        xl_cell.border = THIN_BORDER
    elif style == "summary":
        xl_cell.font = SUMMARY_FONT
        xl_cell.fill = SUMMARY_FILL


def save_workbook(
    path: str,
    month_label: str,
    incomes: list[tuple[str, float]] | None = None,
    expenses: list[tuple[str, float, float]] | None = None,
    currency: str = DEFAULT_CURRENCY,
) -> None:
    wb = build_workbook(month_label, incomes, expenses, currency)
    wb.save(path)
